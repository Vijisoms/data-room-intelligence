"""Generate the mock battery-to-facility crosswalk used to close the PRICED_AT gap.

The sample room names 92 battery codes in the economics and 38 pricing facilities in the basis report but
supplies nothing that connects them, so the catalog relationship PRICED_AT has no links and price answers
are unsupported. This file imitates the crosswalk a seller would supply: one row per battery code with the
facility it sells through. Two battery codes are deliberately unknown to the room (they raise UNMATCHED_KEY).
Run:  .venv/bin/python data/scenarios/crosswalk/make_crosswalk.py
"""

from __future__ import annotations

import random
from pathlib import Path

from openpyxl import Workbook, load_workbook

HERE = Path(__file__).parent
ROOM = HERE.parents[1] / "sample-room"


def battery_codes() -> list[str]:
    workbook = load_workbook(next(ROOM.glob("*ComboCurve*.xlsx")), read_only=True, data_only=True)
    sheet = workbook["ONELINE"]
    rows = sheet.iter_rows(values_only=True)
    header = next(rows)
    column = header.index("BATTERY")
    codes = sorted({row[column] for row in rows if row[column]})
    workbook.close()
    return codes


def facilities() -> list[str]:
    from data_room_intelligence.ingest.families import extract_workbook, load_grids

    path = next(ROOM.rglob("*Basis Report*.xlsx"))
    names: set[str] = set()
    for table in extract_workbook(load_grids(path), path.name, "basis", "basis_report"):
        if table.context.get("family") == "basis":
            names.update(str(cells["C"].value) for _, _, cells in table.records if "C" in cells)
    return sorted(names)


def main() -> None:
    random.seed(20261020)
    codes = battery_codes()
    names = facilities()
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "Crosswalk"
    sheet.append(["Battery Code", "Sales Facility"])
    for index, code in enumerate(codes):
        sheet.append([code, names[index % len(names)] if index < len(names) else random.choice(names)])
    for code in ("BATTERY901", "BATTERY902"):
        sheet.append([code, random.choice(names)])
    sheet.column_dimensions["A"].width = 16
    sheet.column_dimensions["B"].width = 34
    workbook.save(HERE / "Battery Facility Crosswalk_2026.09.xlsx")


if __name__ == "__main__":
    main()
