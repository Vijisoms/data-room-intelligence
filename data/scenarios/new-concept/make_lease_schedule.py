"""Generate the mock lease schedule used to show a new concept entering the model.

Nothing in the base catalog describes a lease, so this file cannot be mapped to an existing type. Its
Lease No. column identifies the rows, its Unit and County columns name units and counties the room already
knows, and the rest are lease terms. Concept discovery proposes a new entity type (Lease) with two
relationships, properties and measures; the data steward reviews and approves it. Two rows name a unit that
does not exist in the room (they raise UNMATCHED_LINK).
Run:  .venv/bin/python data/scenarios/new-concept/make_lease_schedule.py
"""

from __future__ import annotations

import random
from datetime import date, timedelta
from pathlib import Path

from openpyxl import Workbook, load_workbook

HERE = Path(__file__).parent
ROOM = HERE.parents[1] / "sample-room"
LESSORS = ["Cordell Family Trust", "Red Mesa Minerals LP", "Avalon Royalty Partners", "J. T. Whitlock Estate", "Sage Creek Land Co.", "Palomar Ranch Holdings", "Blue Mesa Minerals", "Corona Minerals II", "Harlan & Sons", "Mesquite Flats Trust"]


def units() -> list[tuple[str, str]]:
    workbook = load_workbook(next(ROOM.glob("*Unit Details*.xlsx")), read_only=True, data_only=True)
    sheet = workbook["Unit Details"]
    rows = [(row[2], row[3]) for row in sheet.iter_rows(min_row=8, max_row=97, values_only=True) if row[2]]
    workbook.close()
    return rows


def main() -> None:
    random.seed(20261020)
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "Lease Schedule"
    sheet.append(["Lease No.", "Lessor", "Unit", "County", "Effective Date", "Expiration Date", "Status", "Gross Acres", "Net Acres", "Royalty (%)"])
    number = 1000
    for unit, county in units():
        for _ in range(random.choice([1, 1, 2])):
            number += 1
            gross = round(random.uniform(80, 640), 2)
            effective = date(2019, 1, 1) + timedelta(days=random.randint(0, 2200))
            sheet.append([
                f"LSB-{number}", random.choice(LESSORS), unit, county, effective, effective + timedelta(days=365 * random.choice([3, 5])),
                random.choice(["HBP", "HBP", "Primary Term"]), gross, round(gross * random.uniform(0.35, 1.0), 2), random.choice([18.75, 20.0, 22.5, 25.0]),
            ])
    for unit in ("NORTH FORK EXTENSION", "NORTH FORK EXTENSION"):
        number += 1
        sheet.append([f"LSB-{number}", random.choice(LESSORS), unit, "AGUILA", date(2024, 3, 1), date(2027, 3, 1), "Primary Term", 320.0, 160.0, 25.0])
    for column, width in zip("ABCDEFGHIJ", (12, 28, 30, 14, 14, 15, 13, 12, 12, 12)):
        sheet.column_dimensions[column].width = width
    for row in sheet.iter_rows(min_row=2, min_col=5, max_col=6):
        for cell in row:
            cell.number_format = "yyyy-mm-dd"
    workbook.save(HERE / "Lease Schedule_2026.09.xlsx")


if __name__ == "__main__":
    main()
