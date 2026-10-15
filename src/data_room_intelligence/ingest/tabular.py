from __future__ import annotations

import csv
import io
import re
from dataclasses import dataclass, field
from datetime import date, datetime
from pathlib import Path
from typing import Any

from openpyxl import load_workbook
from openpyxl.utils import column_index_from_string, get_column_letter

from ..contracts import DatasetSchema, FieldSchema

UNIT_PATTERN = re.compile(r"\(([^()]+)\)\s*(?:\.\d+)?$")
UNIT_TOKENS = {"%", "ACRES", "FT", "$", "BBL", "MCF", "BOE", "$/BBL", "$/MCF", "$/BOE"}
MONTH_PATTERN = re.compile(r"^(January|February|March|April|May|June|July|August|September|October|November|December) (\d{4})$")


@dataclass
class Cell:
    value: Any
    ref: str
    row: int
    column: str
    quality: str = "ok"
    is_formula: bool = False


@dataclass
class Table:
    """One extracted table region: records keyed by row, cells keyed by column letter."""

    dataset_id: str
    file: str
    sheet: str
    header_rows: list[int]
    columns: list[str]
    headers: dict[str, list[str]]
    records: list[tuple[str, int, dict[str, Cell]]]
    context: dict[str, Any] = field(default_factory=dict)
    sheet_rows: int = 0
    sheet_columns: int = 0
    formula_cells: int = 0
    uncached_formulas: int = 0

    def header(self, column: str) -> str:
        return " > ".join(self.headers.get(column, [])) or column


@dataclass
class Grid:
    """Cell values and formulas of a worksheet or CSV, 1-based."""

    name: str
    values: list[tuple[Any, ...]]
    formulas: list[tuple[Any, ...]] | None

    @property
    def max_row(self) -> int:
        return len(self.values)

    @property
    def max_col(self) -> int:
        return max((len(row) for row in self.values), default=0)

    def get(self, row: int, col: int) -> Any:
        values = self.values[row - 1] if 0 < row <= len(self.values) else ()
        return values[col - 1] if 0 < col <= len(values) else None

    def formula(self, row: int, col: int) -> str | None:
        if not self.formulas:
            return None
        values = self.formulas[row - 1] if 0 < row <= len(self.formulas) else ()
        text = values[col - 1] if 0 < col <= len(values) else None
        return text if isinstance(text, str) and text.startswith("=") else None


def read_workbook(path: Path) -> list[Grid]:
    values = load_workbook(path, data_only=True, read_only=True)
    formulas = load_workbook(path, data_only=False, read_only=True)
    try:
        return [
            Grid(sheet.title, _trim(list(sheet.iter_rows(values_only=True))), _trim(list(formulas[sheet.title].iter_rows(values_only=True))))
            for sheet in values.worksheets
        ]
    finally:
        values.close()
        formulas.close()


def read_csv(path: Path) -> list[Grid]:
    raw = path.read_bytes()
    for encoding in ("utf-8-sig", "cp1252", "latin-1"):
        try:
            text = raw.decode(encoding)
            break
        except UnicodeDecodeError:
            continue
    try:
        dialect = csv.Sniffer().sniff(text[:8192], delimiters=",;\t|")
    except csv.Error:
        dialect = csv.excel
    rows = [tuple(_parse_text(value) for value in row) for row in csv.reader(io.StringIO(text), dialect)]
    return [Grid(path.stem, _trim(rows), None)]


def _trim(rows: list[tuple[Any, ...]]) -> list[tuple[Any, ...]]:
    while rows and all(value in (None, "") for value in rows[-1]):
        rows.pop()
    return rows


def _parse_text(value: str) -> Any:
    text = value.strip()
    if text == "":
        return None
    try:
        number = float(text.replace(",", "")) if re.fullmatch(r"-?[\d,]*\.?\d+(e-?\d+)?", text, re.I) else None
    except ValueError:
        number = None
    if number is not None:
        return int(number) if number.is_integer() and "." not in text and "e" not in text.lower() else number
    for pattern in ("%Y-%m-%d", "%m/%d/%Y", "%Y-%m-%dT%H:%M:%S", "%Y-%m-%d %H:%M:%S"):
        try:
            return datetime.strptime(text, pattern)
        except ValueError:
            continue
    return text


def detect_header_row(grid: Grid, limit: int = 30) -> int:
    for row in range(1, min(grid.max_row, limit) + 1):
        cells = [grid.get(row, col) for col in range(1, grid.max_col + 1)]
        present = [value for value in cells if value not in (None, "")]
        if len(present) >= 2 and sum(isinstance(value, str) for value in present) / len(present) >= 0.6:
            following = [grid.get(row + 1, col) for col in range(1, grid.max_col + 1)]
            if any(value not in (None, "") for value in following):
                return row
    return 1


def extract_table(
    grid: Grid,
    file: str,
    dataset_id: str,
    header_rows: list[int] | None = None,
    data_rows: tuple[int, int] | None = None,
    key_column: str | None = None,
    fill_right_rows: list[int] | None = None,
    columns: list[str] | None = None,
    stop_at_blank: bool = True,
    row_filter: Any = None,
) -> Table:
    header_rows = header_rows or [detect_header_row(grid)]
    first = (data_rows[0] if data_rows else max(header_rows) + 1)
    last = data_rows[1] if data_rows else grid.max_row
    col_numbers = [column_index_from_string(c) for c in columns] if columns else list(range(1, grid.max_col + 1))
    headers: dict[str, list[str]] = {}
    for hr in header_rows:
        carry = None
        for col in col_numbers:
            value = grid.get(hr, col)
            if isinstance(value, (int, float)) and not isinstance(value, bool):
                continue
            if value in (None, "") and fill_right_rows and hr in fill_right_rows:
                value = carry
            elif value not in (None, ""):
                carry = value
            if value not in (None, ""):
                headers.setdefault(get_column_letter(col), []).append(_header_text(value))
    letters = [get_column_letter(col) for col in col_numbers]
    _dedupe_headers(headers, letters, header_rows)
    records: list[tuple[str, int, dict[str, Cell]]] = []
    blank_run = 0
    for row in range(first, last + 1):
        cells: dict[str, Cell] = {}
        for col, letter in zip(col_numbers, letters):
            value = grid.get(row, col)
            formula = grid.formula(row, col)
            if value in (None, "") and not formula:
                continue
            quality = "uncached_formula" if formula and value is None else "ok"
            cells[letter] = Cell(_json_value(value), f"{letter}{row}", row, letter, quality, formula is not None)
        key_cell = cells.get(key_column) if key_column else None
        if not cells or (key_column and (key_cell is None or key_cell.value in (None, ""))):
            blank_run += 1
            if stop_at_blank and blank_run >= 2 and not data_rows:
                break
            continue
        blank_run = 0
        if row_filter and not row_filter(cells):
            continue
        records.append((str(row), row, cells))
    formula_cells = sum(1 for r in range(1, grid.max_row + 1) for c in range(1, grid.max_col + 1) if grid.formula(r, c))
    uncached = sum(1 for r in range(1, grid.max_row + 1) for c in range(1, grid.max_col + 1) if grid.formula(r, c) and grid.get(r, c) is None)
    return Table(dataset_id, file, grid.name, header_rows, letters, headers, records, {}, grid.max_row, grid.max_col, formula_cells, uncached)


def _header_text(value: Any) -> str:
    if isinstance(value, (datetime, date)):
        return value.strftime("%Y-%m-%d")
    if isinstance(value, float) and value.is_integer():
        return str(int(value))
    return str(value).strip()


def _dedupe_headers(headers: dict[str, list[str]], letters: list[str], header_rows: list[int]) -> None:
    seen: dict[str, int] = {}
    for letter in letters:
        path = headers.get(letter)
        if not path:
            continue
        key = " > ".join(path)
        if key in seen and len(header_rows) == 1:
            seen[key] += 1
            path[-1] = f"{path[-1]}.{seen[key] - 1}"
        else:
            seen.setdefault(key, 1)


def _json_value(value: Any) -> Any:
    if isinstance(value, datetime):
        return value.date().isoformat() if value.time() == datetime.min.time() else value.isoformat()
    if isinstance(value, date):
        return value.isoformat()
    if isinstance(value, str):
        return value.strip() or None
    return value


def profile(table: Table) -> DatasetSchema:
    fields = []
    for letter in table.columns:
        cells = [cells[letter] for _, _, cells in table.records if letter in cells]
        values = [cell.value for cell in cells if cell.value not in (None, "")]
        header_path = table.headers.get(letter, [])
        if not values and not header_path:
            continue
        kind = _infer_type(values)
        unit = _unit(header_path[-1]) if header_path else None
        distinct = {str(value) for value in values}
        fields.append(
            FieldSchema(
                field_id=f"{table.dataset_id}/{letter}",
                dataset_id=table.dataset_id,
                column_ref=letter,
                header_path=header_path or [letter],
                inferred_type=kind,
                unit=unit,
                scale=_scale(unit, values) if kind in {"number", "integer"} else None,
                rows=len(table.records),
                null_rate=round(1 - len(values) / len(table.records), 4) if table.records else 1.0,
                distinct_count=len(distinct),
                samples=values[:3],
                formula_count=sum(cell.is_formula for cell in cells),
                uncached_formula_count=sum(cell.quality == "uncached_formula" for cell in cells),
                candidate_key=bool(table.records) and len(values) == len(table.records) == len(distinct) and kind in {"string", "integer"},
            )
        )
    populated = sum(1 for item in fields if item.inferred_type != "empty")
    return DatasetSchema(
        dataset_id=table.dataset_id,
        file=table.file,
        kind="table",
        locator=f"{table.sheet}!{table.columns[0]}{min(table.header_rows)}:{table.columns[-1]}{table.records[-1][1] if table.records else max(table.header_rows)}",
        header_rows=table.header_rows,
        record_count=len(table.records),
        fields=fields,
        notes=[
            f"sheet used range {table.sheet_rows} rows x {table.sheet_columns} columns",
            f"{len(fields)} headers, {populated} populated",
            f"{table.formula_cells} formulas, {table.uncached_formulas} without cached values",
        ],
    )


def _infer_type(values: list[Any]) -> str:
    if not values:
        return "empty"
    if all(isinstance(value, bool) for value in values):
        return "boolean"
    if all(isinstance(value, (int, float)) and not isinstance(value, bool) for value in values):
        return "integer" if all(float(value).is_integer() for value in values) else "number"
    if all(isinstance(value, str) and re.fullmatch(r"\d{4}-\d{2}-\d{2}", value) for value in values):
        return "date"
    if all(isinstance(value, str) and re.fullmatch(r"\d{4}-\d{2}-\d{2}T[\d:.]+", value) for value in values):
        return "datetime"
    return "string"


def _unit(header: str) -> str | None:
    if header.strip().upper() in UNIT_TOKENS:
        return header.strip()
    match = UNIT_PATTERN.search(header)
    return match.group(1).strip() if match else None


def _scale(unit: str | None, values: list[Any]) -> str | None:
    numbers = [float(value) for value in values if isinstance(value, (int, float))]
    if unit and "%" in unit and numbers:
        return "fraction (0-1)" if max(abs(n) for n in numbers) <= 1.5 else "percent points"
    if unit and unit.upper().startswith("M$"):
        return "thousand dollars"
    return None


def staged_rows(table: Table) -> list[dict[str, Any]]:
    rows = []
    for record_key, row, cells in table.records:
        for letter, cell in cells.items():
            value = cell.value
            rows.append(
                {
                    "dataset_id": table.dataset_id,
                    "record_key": record_key,
                    "row": row,
                    "column": letter,
                    "field_id": f"{table.dataset_id}/{letter}",
                    "num": float(value) if isinstance(value, (int, float)) and not isinstance(value, bool) else None,
                    "text": None if isinstance(value, (int, float)) and not isinstance(value, bool) else (None if value is None else str(value)),
                    "quality": cell.quality,
                }
            )
    return rows


def month_columns(grid: Grid, header_row: int) -> dict[str, str]:
    """Columns whose header is 'Month YYYY' -> ISO month start (wide monthly layouts)."""
    found = {}
    for col in range(1, grid.max_col + 1):
        value = grid.get(header_row, col)
        if isinstance(value, str) and (match := MONTH_PATTERN.match(value.strip())):
            found[get_column_letter(col)] = datetime.strptime(match.group(0), "%B %Y").date().isoformat()
        elif isinstance(value, (datetime, date)):
            found[get_column_letter(col)] = (value.date() if isinstance(value, datetime) else value).replace(day=1).isoformat()
    return found
