"""Declarative layouts for known workbook families; anything else is profiled generically."""

from __future__ import annotations

import re
from datetime import date
from pathlib import Path
from typing import Any

from ..config import slug
from .tabular import Grid, Table, detect_header_row, extract_table, month_columns

BASIS_TAB = re.compile(r"basis report\s*-\s*([a-z]{3})(\d{2})$", re.I)
MONTHS = {name: index for index, name in enumerate(["jan", "feb", "mar", "apr", "may", "jun", "jul", "aug", "sep", "oct", "nov", "dec"], start=1)}


# Headers a sheet must carry before a built-in recipe is trusted with it. A file whose name resembles
# the sample room's but whose layout differs is profiled generically instead (and mapped in the Model step).
ONELINE_REQUIRED = {"INPT ID", "API 14", "Well Name", "UNIT", "Current Operator", "Project Name"}
CASHFLOW_REQUIRED = {"Date", "Group"}


def _looks_like_unit_details(grid: Grid) -> bool:
    header = grid.get(7, 3)
    return grid.max_row >= 8 and isinstance(header, str) and header.strip().upper() == "UNIT"


def extract_workbook(grids: list[Grid], file: str, file_id: str, role: str | None) -> list[Table]:
    tables: list[Table] = []
    for grid in grids:
        dataset = f"{file_id}/{slug(grid.name)}"
        if role == "unit_details" and grid.name == "Unit Details" and _looks_like_unit_details(grid):
            table = extract_table(grid, file, dataset, header_rows=[5, 6, 7], fill_right_rows=[5], data_rows=(8, grid.max_row), key_column="C")
            table.context["family"] = "unit_details"
            tables.append(table)
        elif role == "lease_operating_statement" and month_columns(grid, 2):
            tables.append(_los(grid, file, dataset))
        elif role == "basis_report" and BASIS_TAB.search(grid.name):
            tables.extend(_basis(grid, file, dataset))
        elif grid.max_row and grid.max_col:
            table = extract_table(grid, file, dataset, header_rows=[detect_header_row(grid)])
            names = {path[-1] for path in table.headers.values()}
            family = "generic"
            if role == "economics_oneline" and grid.name == "ONELINE" and ONELINE_REQUIRED <= names:
                family = "economics_oneline"
            elif role == "economics_oneline" and grid.name == "CASHFLOW" and CASHFLOW_REQUIRED <= names:
                family = "economics_cashflow"
            table.context["family"] = family
            tables.append(table)
    return tables


def _los(grid: Grid, file: str, dataset: str) -> Table:
    months = month_columns(grid, 2)
    table = extract_table(grid, file, dataset, header_rows=[2], data_rows=(3, grid.max_row), key_column="A", columns=["A", "B", *months])
    section = None
    sections: dict[int, str | None] = {}
    for _, row, cells in table.records:
        label = str(cells["A"].value)
        has_values = any(letter in months and (cells[letter].value is not None or cells[letter].quality != "ok") for letter in cells)
        if label.endswith(":") or not has_values:
            section = label.rstrip(":")
        sections[row] = section
    table.context.update({"family": "los", "scope": grid.name, "title": grid.get(1, 1), "months": months, "sections": sections})
    return table


def _basis(grid: Grid, file: str, dataset: str) -> list[Table]:
    match = BASIS_TAB.search(grid.name)
    period = date(2000 + int(match.group(2)), MONTHS[match.group(1).lower()], 1).isoformat()
    columns = [_letter(index) for index in range(3, grid.max_col + 1)]

    def facility(cells: dict[str, Any]) -> bool:
        name = cells.get("C")
        numbers = sum(isinstance(cell.value, (int, float)) for cell in cells.values())
        return (
            name is not None
            and isinstance(name.value, str)
            and name.value != "Month"
            and not re.fullmatch(r"\d{4}-\d{2}-\d{2}", name.value)
            and numbers >= 3
        )

    previous_end = _find(grid, "C", "Forward 12 Months", start=10) or grid.max_row
    previous = extract_table(grid, file, f"{dataset}/previous-month", header_rows=[8, 9], fill_right_rows=[8], data_rows=(10, previous_end - 1), key_column="C", columns=columns, row_filter=facility)
    forward = extract_table(grid, file, f"{dataset}/forward-12-months", header_rows=[previous_end, 9], fill_right_rows=[previous_end], data_rows=(previous_end + 1, grid.max_row), key_column="C", columns=columns, row_filter=facility)
    strip_header = _find(grid, "C", "Month", start=previous_end)
    tables = [previous, forward]
    if strip_header:
        strip = extract_table(grid, file, f"{dataset}/strip-prices", header_rows=[strip_header - 1, strip_header], fill_right_rows=[strip_header - 1], data_rows=(strip_header + 1, grid.max_row), key_column="C", columns=columns[:8], row_filter=lambda cells: isinstance(cells["C"].value, str) and bool(re.fullmatch(r"\d{4}-\d{2}-\d{2}", cells["C"].value)))
        strip.context.update({"family": "basis_strip", "period": period})
        tables.append(strip)
    for table, section in ((previous, "previous_month"), (forward, "forward_12_months")):
        table.context.update({"family": "basis", "section": section, "period": period})
    return tables


def _find(grid: Grid, column: str, text: str, start: int = 1) -> int | None:
    col = ord(column) - 64
    for row in range(start, grid.max_row + 1):
        value = grid.get(row, col)
        if isinstance(value, str) and value.strip() == text:
            return row
    return None


def _letter(index: int) -> str:
    from openpyxl.utils import get_column_letter

    return get_column_letter(index)


def load_grids(path: Path) -> list[Grid]:
    from .tabular import read_csv, read_workbook

    return read_csv(path) if path.suffix.lower() in {".csv", ".tsv", ".txt"} else read_workbook(path)
