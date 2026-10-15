from __future__ import annotations

import re
from collections import Counter
from collections.abc import Iterator
from contextlib import contextmanager
from contextvars import ContextVar
from datetime import date, datetime
from pathlib import Path
from typing import Any

import pandas as pd
import shapefile
from openpyxl import load_workbook
from openpyxl.utils import get_column_letter
from pypdf import PdfReader
from pyproj import Geod
from shapely.geometry import shape
from shapely.ops import unary_union

from .models import Metric, MonthlyProduction, SourceRef, UnitRecord, WellRecord

SELLER_OPERATOR = "LONE STAR BASIN OPERATING LLC"
PV10_COLUMN = "10% Discount Cash Flow (M$)"
PV10_ALT_COLUMN = "10% Discount Cash Flow (M$).1"
_ROOM_ROOT: ContextVar[Path | None] = ContextVar("data_room_root", default=None)

ZONES = [
    "USS",
    "MSS",
    "JM",
    "LSS",
    "DEAN",
    "WCA",
    "WCB I",
    "WCB II",
    "WCC",
    "WCD",
    "STR / BND",
    "UP BRNT",
    "LR BRNT",
    "MSSP",
    "WDFD",
]


def source(path: Path, locator: str, as_of: str | None = None) -> SourceRef:
    return SourceRef(path=room_path(path), locator=locator, as_of=as_of)


@contextmanager
def data_room_root(root: Path) -> Iterator[None]:
    """Cite parsed sources relative to this data room instead of the local disk."""
    token = _ROOM_ROOT.set(root.resolve())
    try:
        yield
    finally:
        _ROOM_ROOT.reset(token)


def room_path(path: Path) -> str:
    root = _ROOM_ROOT.get()
    if root is not None:
        try:
            return path.resolve().relative_to(root).as_posix()
        except ValueError:
            pass
    return path.name


def parse_unit_details(path: Path) -> tuple[list[UnitRecord], list[Metric]]:
    workbook = load_workbook(path, data_only=True, read_only=False)
    sheet = workbook["Unit Details"]
    units: list[UnitRecord] = []
    for row_number, values in enumerate(
        sheet.iter_rows(min_row=8, max_row=97, min_col=1, max_col=137, values_only=True),
        start=8,
    ):
        if not values[2]:
            continue
        zone_net_acres = {zone: float(values[10 + index] or 0) for index, zone in enumerate(ZONES)}
        zone_nri = {zone: float(values[26 + index] or 0) for index, zone in enumerate(ZONES)}
        zone_wi = {zone: float(values[90 + index] or 0) for index, zone in enumerate(ZONES)}
        zone_locations = {
            zone: float(values[122 + index] or 0) for index, zone in enumerate(ZONES)
        }
        name = str(values[2]).strip()
        units.append(
            UnitRecord(
                unit_key=_key(name),
                unit_name=name,
                county=str(values[3]).strip(),
                development_area=str(values[5]).strip(),
                gross_acres=float(values[7]),
                lateral_length_ft=float(values[8]),
                zone_net_acres=zone_net_acres,
                zone_wi=zone_wi,
                zone_nri=zone_nri,
                zone_gross_locations=zone_locations,
                source=source(path, f"Unit Details!C{row_number}:EG{row_number}", "2026-10-01"),
            )
        )
    workbook.close()
    gross_acres = sum(item.gross_acres for item in units)
    metrics = [
        Metric(
            name="unit_count",
            value=len(units),
            unit="count",
            sources=[source(path, "Unit Details!C8:C97", "2026-10-01")],
        ),
        Metric(
            name="gross_acres",
            value=gross_acres,
            unit="acres",
            sources=[source(path, "Unit Details!H8:H97", "2026-10-01")],
            method="sum of contractual unit acres",
        ),
        Metric(
            name="county_count",
            value=len({item.county for item in units}),
            unit="count",
            sources=[source(path, "Unit Details!D8:D97", "2026-10-01")],
        ),
        Metric(
            name="development_area_count",
            value=len({item.development_area for item in units}),
            unit="count",
            sources=[source(path, "Unit Details!F8:F97", "2026-10-01")],
        ),
    ]
    for zone in ZONES:
        metrics.extend(
            [
                Metric(
                    name="net_acres",
                    value=sum(item.zone_net_acres[zone] for item in units),
                    unit="acres",
                    dimensions={"zone": zone},
                    sources=[source(path, f"Unit Details zone {zone}", "2026-10-01")],
                    method="sum of unit-level net acres",
                ),
                Metric(
                    name="net_locations",
                    value=sum(
                        item.zone_gross_locations[zone] * item.zone_wi[zone] for item in units
                    ),
                    unit="net wells",
                    dimensions={"zone": zone},
                    sources=[source(path, f"Unit Details zone {zone}", "2026-10-01")],
                    method="sum of gross locations multiplied by unit WI",
                ),
            ]
        )
    return units, metrics


def parse_combocurve(path: Path) -> tuple[list[WellRecord], list[dict[str, Any]], list[Metric]]:
    workbook = load_workbook(path, data_only=True, read_only=True)
    try:
        header, oneline = _sheet_records(workbook["ONELINE"])
        _, cash_flow = _sheet_records(workbook["CASHFLOW"])
    finally:
        workbook.close()
    wells: list[WellRecord] = []
    for row_number, row in oneline:
        api_14 = _api_number(row.get("API 14"), 14)
        chosen = _optional_text(row.get("Chosen ID"))
        name = _optional_text(row.get("Well Name")) or f"well-{row_number - 1}"
        if api_14:
            well_key = api_14
            well_key_source = "api_14"
        elif chosen:
            well_key = chosen
            well_key_source = "chosen_id"
        else:
            well_key = _key(name)
            well_key_source = "well_name"
        wells.append(
            WellRecord(
                well_key=well_key,
                chosen_id=chosen,
                well_key_source=well_key_source,
                api_10=_api_number(row.get("API 10"), 10),
                api_14=api_14,
                well_name=name,
                unit_key=_key(_optional_text(row.get("UNIT")) or "") or None,
                operator=_optional_text(row.get("Current Operator")),
                operator_alias=_optional_text(row.get("Current Operator Alias")),
                status=(_optional_text(row.get("WELL STATUS")) or "UNKNOWN").upper(),
                landing_zone=_optional_text(row.get("Landing Zone")),
                county=_optional_text(row.get("County/Parish")),
                wi=_optional_float(row.get("WI Oil (%)")),
                nri=_optional_float(row.get("NRI Oil (%)")),
                net_boe_mboe=_optional_float(row.get("Net BOE Sales Volume (MBBL)")),
                pv10_musd=_thousand_dollars_to_millions(row.get(PV10_COLUMN)),
                pv10_alt_musd=_thousand_dollars_to_millions(row.get(PV10_ALT_COLUMN)),
                irr=_percent_to_fraction(row.get("IRR (%)")),
                source=source(path, f"ONELINE!row {row_number}", "2026-10-01"),
            )
        )
    cash_records = [record for _, record in cash_flow]
    last_row = oneline[-1][0] if oneline else 2

    def column(name: str) -> str:
        return _column_range(header, name, "ONELINE", 2, last_row)

    counts = Counter(item.status for item in wells)
    metrics = [
        Metric(
            name="economic_location_count",
            value=len(wells),
            unit="count",
            sources=[source(path, "SUMMARY!B10", "2026-10-01")],
        ),
        Metric(
            name="producing_well_count",
            value=counts["PRODUCING"],
            unit="count",
            sources=[source(path, "SUMMARY!B8", "2026-10-01")],
        ),
        Metric(
            name="average_wi",
            value=sum(item.wi or 0 for item in wells) / len(wells),
            unit="fraction",
            sources=[source(path, column("WI Oil (%)"), "2026-10-01")],
        ),
        Metric(
            name="average_nri",
            value=sum(item.nri or 0 for item in wells) / len(wells),
            unit="fraction",
            sources=[source(path, column("NRI Oil (%)"), "2026-10-01")],
        ),
        Metric(
            name="operated_location_fraction",
            value=sum(item.operator == SELLER_OPERATOR for item in wells) / len(wells),
            unit="fraction",
            sources=[source(path, column("Current Operator"), "2026-10-01")],
            method=(
                f"share of locations whose Current Operator is {SELLER_OPERATOR}; "
                "definition unconfirmed (see OPERATOR_ALIAS_CONFLICT)"
            ),
        ),
    ]
    return wells, cash_records, metrics


def parse_gis(
    unit_path: Path, leasehold_path: Path, pdp_path: Path
) -> tuple[list[dict[str, Any]], list[Metric]]:
    geod = Geod(ellps="WGS84")

    def read(path: Path) -> tuple[list[dict[str, Any]], list[Any]]:
        reader = shapefile.Reader(str(path))
        fields = [field[0] for field in reader.fields[1:]]
        records: list[dict[str, Any]] = []
        geometries: list[Any] = []
        for shape_record in reader.iterShapeRecords():
            records.append(dict(zip(fields, list(shape_record.record))))
            geometries.append(shape(shape_record.shape.__geo_interface__))
        return records, geometries

    unit_records, unit_geometries = read(unit_path)
    _, lease_geometries = read(leasehold_path)
    pdp_records, pdp_geometries = read(pdp_path)
    unit_union = unary_union(unit_geometries)
    lease_union = unary_union(lease_geometries)

    def geodesic_area(geometry: Any) -> float:
        area, _ = geod.geometry_area_perimeter(geometry)
        return abs(area) / 4046.8564224

    links: list[dict[str, Any]] = []
    for record, lateral in zip(pdp_records, pdp_geometries):
        candidates = [
            (unit_record, geometry.intersection(lateral).length)
            for unit_record, geometry in zip(unit_records, unit_geometries)
            if geometry.intersects(lateral)
        ]
        matched = max(candidates, key=lambda item: item[1])[0] if candidates else None
        links.append(
            {
                "api_10": _api_number(record.get("API_10"), 10),
                "api_14": _api_number(record.get("API_14"), 14),
                "well_name": record.get("Well_Name"),
                "landing_zone": record.get("Landing_Zo"),
                "unit_name": matched.get("UNIT_NAME") if matched else None,
                "unit_key": _key(str(matched.get("UNIT_NAME"))) if matched else None,
                "intersects_leasehold": lease_union.intersects(lateral),
            }
        )
    unit_area = geodesic_area(unit_union)
    lease_area = geodesic_area(lease_union)
    metrics = [
        Metric(
            name="mapped_unit_area",
            value=unit_area,
            unit="acres",
            sources=[source(unit_path, "dissolved polygon geodesic area")],
            method="WGS84 geodesic area of dissolved polygons",
        ),
        Metric(
            name="active_leasehold_area",
            value=lease_area,
            unit="acres",
            sources=[source(leasehold_path, "dissolved polygon geodesic area")],
            method="WGS84 geodesic area of dissolved polygons",
        ),
        Metric(
            name="pdp_lateral_count",
            value=len(pdp_geometries),
            unit="count",
            sources=[source(pdp_path, "all features")],
        ),
        Metric(
            name="pdp_laterals_intersecting_leasehold",
            value=sum(item["intersects_leasehold"] for item in links),
            unit="count",
            sources=[source(pdp_path, "all features"), source(leasehold_path, "all features")],
            method="spatial intersection",
        ),
    ]
    return links, metrics


def parse_basis(path: Path) -> list[dict[str, Any]]:
    workbook = load_workbook(path, data_only=True, read_only=False)
    results: list[dict[str, Any]] = []
    for sheet in workbook.worksheets[1:]:
        rows: list[tuple[Any, ...]] = []
        for row in sheet.iter_rows(
            min_row=10, max_row=sheet.max_row, min_col=3, max_col=28, values_only=True
        ):
            if row[0] and row[2] == "HORZ" and isinstance(row[7], (int, float)):
                rows.append(row)
        if not rows:
            continue

        def average(position: int) -> float | None:
            values = [float(row[position]) for row in rows if isinstance(row[position], (int, float))]
            return sum(values) / len(values) if values else None

        results.append(
            {
                "period": sheet.title,
                "facility_count": len(rows),
                "oil_nymex_usd_bbl": average(3),
                "oil_basis_usd_bbl": average(4),
                "net_oil_usd_bbl": average(7),
                "gas_hh_usd_mcf": average(11),
                "gas_basis_usd_mcf": average(12),
                "net_gas_usd_mcf": average(15),
                "ngl_usd_gal": average(20),
                "source": source(path, f"{sheet.title}!C10:AB{sheet.max_row}").model_dump(),
            }
        )
    workbook.close()
    return results


def parse_los(path: Path) -> list[MonthlyProduction]:
    workbook = load_workbook(path, data_only=True, read_only=False)
    production: list[MonthlyProduction] = []
    for sheet in workbook.worksheets:
        label_rows = {
            str(sheet.cell(row, 1).value).strip(): row
            for row in range(1, sheet.max_row + 1)
            if sheet.cell(row, 1).value not in (None, "")
        }
        loe_row = label_rows.get("Total LOE")
        income_row = label_rows.get("Operating Income (Loss)")
        for column in range(1, sheet.max_column + 1):
            raw = sheet.cell(2, column).value
            if not isinstance(raw, str):
                continue
            try:
                month = datetime.strptime(raw.strip(), "%B %Y").date()
            except ValueError:
                continue
            values = [sheet.cell(row, column).value for row in range(3, 7)]
            if not all(isinstance(value, (int, float)) for value in values):
                continue
            revenue = sheet.cell(12, column).value
            loe = sheet.cell(loe_row, column).value if loe_row else None
            operating_income = sheet.cell(income_row, column).value if income_row else None
            production.append(
                MonthlyProduction(
                    asset_segment=sheet.title,
                    month=month.isoformat(),
                    oil_bbl=float(values[0]),
                    gas_mcf=float(values[1]),
                    ngl_bbl=float(values[2]),
                    boe=float(values[3]),
                    revenue_usd=float(revenue) if isinstance(revenue, (int, float)) else None,
                    loe_usd=float(loe) if isinstance(loe, (int, float)) else None,
                    operating_income_usd=(
                        float(operating_income)
                        if isinstance(operating_income, (int, float))
                        else None
                    ),
                    source=source(path, f"{sheet.title}!{sheet.cell(2, column).coordinate}:{sheet.cell(12, column).coordinate}"),
                )
            )
    workbook.close()
    return production


def parse_pdfs(paths: list[Path]) -> dict[str, Any]:
    facts: dict[str, Any] = {}
    for path in paths:
        reader = PdfReader(str(path))
        text = "\n".join(page.extract_text() or "" for page in reader.pages)
        facts[path.name] = {
            "page_count": len(reader.pages),
            "text": text,
            "metrics": _pdf_metrics(text),
            "source": source(path, "all pages").model_dump(),
        }
    return facts


def _pdf_metrics(text: str) -> dict[str, float]:
    patterns = {
        "drilling_units": r"Drilling Units\s+([\d,]+)",
        "gross_acres": r"Gross Acres\s+([\d,]+)",
        "producing_wells": r"Producing Wells \(PDP\)\s+([\d,]+)",
        "economic_locations": r"Total Econ Locations\s+([\d,]+)",
        "dc_capital_musd": r"D&C Capital \(\$MM\)\s+\$?([\d.]+)",
        "eur_mboe": r"EUR \(MBOE\)\s+([\d,]+)",
        "pv10_musd": r"PV-10 \(\$MM\)\s+\$?([\d.]+)",
        "irr_pct": r"IRR \(%\)\s+([\d.]+)",
        "payout_months": r"Payout \(months\)\s+([\d.]+)",
        "breakeven_usd_bbl": r"Breakeven \(\$/Bbl WTI\)\s+\$?([\d.]+)",
    }
    result: dict[str, float] = {}
    for key, pattern in patterns.items():
        match = re.search(pattern, text)
        if match:
            result[key] = float(match.group(1).replace(",", ""))
    return result


def _key(value: str) -> str:
    return re.sub(r"[^A-Z0-9]+", "_", value.upper()).strip("_")


def _sheet_records(sheet: Any) -> tuple[list[str], list[tuple[int, dict[str, Any]]]]:
    """Read a header-row worksheet as (row number, record) pairs, skipping blank rows."""
    rows = sheet.iter_rows(values_only=True)
    header = _unique_headers(next(rows, ()))
    records: list[tuple[int, dict[str, Any]]] = []
    for row_number, values in enumerate(rows, start=2):
        if all(value in (None, "") for value in values):
            continue
        records.append(
            (row_number, {name: _json_value(value) for name, value in zip(header, values)})
        )
    return header, records


def _unique_headers(values: tuple[Any, ...]) -> list[str]:
    seen: Counter[str] = Counter()
    headers: list[str] = []
    for index, value in enumerate(values, start=1):
        name = _optional_text(value) or f"column_{get_column_letter(index)}"
        seen[name] += 1
        headers.append(name if seen[name] == 1 else f"{name}.{seen[name] - 1}")
    return headers


def _column_range(header: list[str], name: str, sheet: str, first: int, last: int) -> str:
    letter = get_column_letter(header.index(name) + 1)
    return f"{sheet}!{letter}{first}:{letter}{last}"


def _json_value(value: Any) -> Any:
    return value.isoformat() if isinstance(value, (datetime, date)) else value


def _api_number(value: Any, digits: int) -> str | None:
    """Return an API well number as zero-padded digit text (Excel stores it as a number)."""
    if isinstance(value, float) and value.is_integer():
        value = int(value)
    text = _optional_text(value)
    if text is None:
        return None
    return text.zfill(digits) if text.isdigit() else text


def _optional_text(value: Any) -> str | None:
    if value is None or (isinstance(value, float) and pd.isna(value)):
        return None
    text = str(value).strip()
    return text or None


def _optional_float(value: Any) -> float | None:
    try:
        result = float(value)
        return None if pd.isna(result) else result
    except (TypeError, ValueError):
        return None


def _thousand_dollars_to_millions(value: Any) -> float | None:
    parsed = _optional_float(value)
    return parsed / 1000 if parsed is not None else None


def _percent_to_fraction(value: Any) -> float | None:
    parsed = _optional_float(value)
    return parsed / 100 if parsed is not None else None
