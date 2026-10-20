"""FastAPI service: projects, graph exploration, file viewers, model changes, review, mappings, uploads and chat.

Every data endpoint is scoped to a project with the `room` query parameter (default: the sample room from
DATA_ROOM_ASSET_ID). A project is a data room with its own files, snapshots, decisions and model extension.
"""

from __future__ import annotations

import csv
import io
import json
import os
import re
import threading
import zipfile
from collections import Counter, defaultdict
from functools import lru_cache
from pathlib import Path
from typing import Any

from fastapi import Depends, FastAPI, File, HTTPException, Query, UploadFile
from fastapi.responses import FileResponse, JSONResponse, StreamingResponse
from fastapi.staticfiles import StaticFiles
from openpyxl.utils import column_index_from_string
from pydantic import BaseModel, Field

from ..config import REPO_ROOT, RoomPaths, offline, slug
from ..ingest.families import load_grids
from ..ontology.catalog import catalog_changes, load_aliases, load_rules, room_catalog
from ..ontology.diff import snapshot_diff
from ..ontology.mapping import header_key
from ..ontology.modeling import decide_proposal, get_proposal, list_proposals
from ..ontology.pipeline import publish
from ..ontology.projects import DEFAULT_ROOM_ID, TEMPLATES, create_project, delete_project, get_project, list_projects, room_dir
from ..ontology.reviews import issue_decisions, record_issue_decision, record_mapping_decision
from ..ontology.store import Snapshot, list_rooms, list_snapshots
from . import chat, operations as ops

UPLOAD_TYPES = {".xlsx", ".xlsm", ".csv", ".pdf", ".zip"}
ZIP_MEMBER_TYPES = {".xlsx", ".xlsm", ".csv", ".pdf", ".shp", ".shx", ".dbf", ".prj", ".cpg", ".sbn", ".sbx", ".qix"}
MAX_UPLOAD_BYTES = 200 * 1024 * 1024


class AskRequest(BaseModel):
    question: str = Field(min_length=1, max_length=500)
    scope: str | None = None


class IssueDecision(BaseModel):
    issue_id: str | None = None
    rule: str
    scope: str = Field(default="issue", pattern="^(issue|rule)$")
    choice: str = Field(min_length=1, max_length=120)
    reason: str = Field(min_length=1, max_length=1000)
    reviewer: str = Field(default="reviewer", max_length=80)


class MappingDecision(BaseModel):
    field_id: str
    target: str
    status: str = Field(pattern="^(approved|rejected)$")
    reviewer: str = Field(default="reviewer", max_length=80)
    reason: str = Field(default="", max_length=1000)


class ProposalDecision(BaseModel):
    status: str = Field(pattern="^(approved|rejected)$")
    reviewer: str = Field(default="data steward", max_length=80)
    reason: str = Field(default="", max_length=1000)
    items: list[dict[str, Any]] | None = Field(default=None, description="Edited change set; defaults to the proposal's items")
    rebuild: bool = True


class PublishRequest(BaseModel):
    reviewer: str | None = None
    reason: str | None = None
    force: bool = False


class NewProject(BaseModel):
    name: str = Field(min_length=1, max_length=80)
    description: str = Field(default="", max_length=500)
    template: str = Field(default="base", pattern="^(base|copy)$")
    copied_from: str | None = None


class State:
    def __init__(self) -> None:
        self.lock = threading.Lock()
        self._snapshots: dict[str, Snapshot] = {}

    def snapshot(self, room_id: str) -> Snapshot:
        with self.lock:
            snap = self._snapshots.get(room_id)
            if snap is None:
                if get_project(room_id) is None:
                    raise HTTPException(404, f"unknown project {room_id}")
                try:
                    snap = self._snapshots[room_id] = Snapshot(room_id)
                except FileNotFoundError as error:
                    raise HTTPException(503, f"No snapshot built yet for {room_id}: add files or run `data-room-intel snapshot --room-id {room_id}` ({error})") from error
            return snap

    def reload(self, room_id: str) -> Snapshot:
        with self.lock:
            self._snapshots.pop(room_id, None)
        _grids.cache_clear()
        return self.snapshot(room_id)


state = State()
app = FastAPI(title="Data Room Intelligence", version="0.2.0")


def room_param(room: str | None = Query(None, description="Project (room) ID; defaults to the sample room")) -> str:
    return room or DEFAULT_ROOM_ID


def current(room: str = Depends(room_param)) -> Snapshot:
    return state.snapshot(room)


@lru_cache(maxsize=8)
def _grids(room_id: str, snapshot_id: str, file_id: str) -> dict[str, Any]:
    snap = state.snapshot(room_id)
    entry = _entry(snap, file_id)
    return {grid.name: grid for grid in load_grids(snap.dir / "bronze" / "originals" / entry["path"])}


def _entry(snap: Snapshot, file_id: str) -> dict[str, Any]:
    entry = next((f for f in snap.manifest["files"] if f["file_id"] == file_id), None)
    if entry is None:
        raise HTTPException(404, f"unknown file {file_id}")
    return entry


def _issue_counts(snap: Snapshot, entity: str) -> dict[str, int]:
    return dict(Counter(i["severity"] for i in snap.issues_by_subject.get(entity, []) if i["status"] == "open"))


def _conflict_edges(snap: Snapshot) -> set[str]:
    edges = set()
    for issue in snap.issues:
        if issue["rule"] != "UNIT_ASSIGNMENT_CONFLICT" or issue["status"] != "open":
            continue
        well = issue["subject"]
        edges.update(e["id"] for e in snap.out.get(well, []) if e["kind"] == "LATERAL_IN_UNIT")
        for link in (e for e in snap.inn.get(well, []) if e["kind"] == "IS_WELL"):
            edges.update(e["id"] for e in snap.out.get(link["source"], []) if e["kind"] == "ASSIGNED_TO_UNIT")
    return edges


def _node_card(snap: Snapshot, node: dict[str, Any], level: int | None = None) -> dict[str, Any]:
    props = {k: v for k, v in node["props"].items() if k not in {"raw"}}
    return {"id": node["id"], "type": node["type"], "label": node["label"], "level": level, "issues": _issue_counts(snap, node["id"]), "props": props}


def _type_index(snap: Snapshot, type_: str) -> int:
    order = snap.type_order
    return order.index(type_) if type_ in order else len(order)


# ---------------------------------------------------------------------------
# Health and projects
# ---------------------------------------------------------------------------


@app.get("/api/health")
def health() -> dict[str, Any]:
    return {"status": "ok", "offline": offline(), "model": None if offline() else os.environ.get("FOUNDRY_MODEL"), "room_id": DEFAULT_ROOM_ID}


def _project_summary(record: dict[str, Any]) -> dict[str, Any]:
    room_id = record["room_id"]
    summary: dict[str, Any] = {**record, "has_snapshot": False, "status": "empty", "files": 0, "open_issues": {}, "blocking_questions": 0, "decided_issues": 0, "pending_mappings": 0, "pending_proposals": 0, "model_changes": len(catalog_changes(room_id)), "catalog_version": room_catalog(room_id).catalog_version, "built_at": None, "snapshot_id": None, "counts": None}
    try:
        snap = state.snapshot(room_id)
    except HTTPException:
        incoming = RoomPaths(room_id).incoming
        folder = room_dir(record)
        summary["files"] = sum(1 for p in [*(folder.rglob("*") if folder.is_dir() else []), *(incoming.rglob("*") if incoming.is_dir() else [])] if p.is_file() and not p.name.startswith((".", "~$")))
        summary["pending_proposals"] = sum(p["status"] == "proposed" for p in list_proposals(room_id))
        return summary
    open_issues = [i for i in snap.issues if i["status"] == "open"]
    summary.update(
        {
            "has_snapshot": True,
            "status": snap.meta["status"],
            "snapshot_id": snap.snapshot_id,
            "built_at": snap.meta.get("built_at"),
            "catalog_version": snap.meta.get("catalog_version"),
            "files": sum(1 for f in snap.manifest["files"] if f["kind"] != "gis_sidecar"),
            "counts": snap.meta["counts"],
            "open_issues": dict(Counter(i["severity"] for i in open_issues)),
            "blocking_questions": len({i["rule"] for i in open_issues if i["severity"] == "blocking"}),
            "decided_issues": len(snap.issues) - len(open_issues),
            "pending_mappings": sum(m["status"] == "proposed" for m in snap.mappings),
            "pending_proposals": sum(p["status"] == "proposed" for p in list_proposals(room_id)),
            "asset_name": next((n["label"] for n in snap.nodes.values() if n["type"] == "Asset"), record["name"]),
        }
    )
    return summary


@app.get("/api/projects")
def projects() -> dict[str, Any]:
    return {"projects": [_project_summary(r) for r in list_projects()], "templates": TEMPLATES, "default": DEFAULT_ROOM_ID}


@app.post("/api/projects")
def new_project(request: NewProject) -> dict[str, Any]:
    try:
        record = create_project(request.name, request.description, request.template, request.copied_from)
    except ValueError as error:
        raise HTTPException(400, str(error)) from error
    return _project_summary(record)


@app.get("/api/projects/{room_id}")
def project(room_id: str) -> dict[str, Any]:
    record = get_project(room_id)
    if record is None:
        raise HTTPException(404, f"unknown project {room_id}")
    return _project_summary(record)


@app.delete("/api/projects/{room_id}")
def remove_project(room_id: str) -> dict[str, Any]:
    if get_project(room_id) is None:
        raise HTTPException(404, f"unknown project {room_id}")
    try:
        delete_project(room_id)
    except ValueError as error:
        raise HTTPException(400, str(error)) from error
    with state.lock:
        state._snapshots.pop(room_id, None)
    return {"deleted": room_id}


@app.get("/api/rooms")
def rooms() -> list[dict[str, Any]]:
    return list_rooms()


@app.get("/api/room")
def room(snap: Snapshot = Depends(current)) -> dict[str, Any]:
    open_issues = [i for i in snap.issues if i["status"] == "open"]
    record = get_project(snap.room_id) or {}
    return {
        **{k: snap.meta.get(k) for k in ("room_id", "snapshot_id", "status", "catalog_version", "rules_version", "built_at", "mappings", "model_proposals")},
        "files": sum(1 for f in snap.manifest["files"] if f["kind"] != "gis_sidecar"),
        "name": record.get("name") or next((n["label"] for n in snap.nodes.values() if n["type"] == "Asset"), snap.room_id),
        "counts": snap.meta["counts"],
        "open_issues": dict(Counter(i["severity"] for i in open_issues)),
        "decided_issues": len(snap.issues) - len(open_issues),
        "pending_proposals": sum(p["status"] == "proposed" for p in list_proposals(snap.room_id)),
        "policies": {rule: snap.policy(rule) for rule in {i["rule"] for i in snap.issues} if snap.policy(rule)},
    }


# ---------------------------------------------------------------------------
# Catalog, model proposals and change log
# ---------------------------------------------------------------------------


@app.get("/api/catalog")
def catalog(snap: Snapshot = Depends(current)) -> dict[str, Any]:
    kinds = Counter(e["kind"] for e in snap.edges)
    types = Counter(n["type"] for n in snap.nodes.values())
    rules = load_rules()
    base_types = {t.name for t in room_catalog(None).entity_types}
    base_rels = {r.name for r in room_catalog(None).relationship_types}
    base_measures = {(m.subject, m.name) for m in room_catalog(None).measures}
    data = snap.catalog.model_dump(mode="json")
    for entity in data["entity_types"]:
        entity["origin"] = "base" if entity["name"] in base_types else "project"
    for rel in data["relationship_types"]:
        rel["origin"] = "base" if rel["name"] in base_rels else "project"
    for measure in data["measures"]:
        measure["origin"] = "base" if (measure["subject"], measure["name"]) in base_measures else "project"
    return {
        **data,
        "rules": [{**r.model_dump(mode="json"), "count": sum(i["rule"] == r.code for i in snap.issues)} for r in rules.rules],
        "aliases": {name: table.model_dump(mode="json") for name, table in load_aliases().items()},
        "counts": {"entities": dict(types), "relationships": dict(kinds)},
        "changes": catalog_changes(snap.room_id),
    }


@app.get("/api/model/proposals")
def model_proposals(room: str = Depends(room_param), status: str | None = None) -> list[dict[str, Any]]:
    if get_project(room) is None:
        raise HTTPException(404, f"unknown project {room}")
    return [p for p in list_proposals(room) if not status or p["status"] == status]


@app.get("/api/model/proposals/{proposal_id}")
def model_proposal(proposal_id: str, room: str = Depends(room_param)) -> dict[str, Any]:
    proposal = get_proposal(room, proposal_id)
    if proposal is None:
        raise HTTPException(404, "unknown proposal")
    return proposal


@app.post("/api/model/proposals/{proposal_id}/decision")
async def model_proposal_decision(proposal_id: str, decision: ProposalDecision, room: str = Depends(room_param)) -> dict[str, Any]:
    try:
        proposal = decide_proposal(room, proposal_id, decision.status, decision.reviewer, decision.reason, decision.items)
    except KeyError as error:
        raise HTTPException(404, "unknown proposal") from error
    except ValueError as error:
        raise HTTPException(400, f"the change set is not valid: {error}") from error
    result: dict[str, Any] = {"proposal": proposal}
    if decision.status == "approved" and decision.rebuild:
        result["rebuild"] = await rebuild(room)
    return result


@app.get("/api/model/changes")
def model_changes(room: str = Depends(room_param)) -> list[dict[str, Any]]:
    if get_project(room) is None:
        raise HTTPException(404, f"unknown project {room}")
    return catalog_changes(room)


# ---------------------------------------------------------------------------
# Graph
# ---------------------------------------------------------------------------


@app.get("/api/graph/stats")
def graph_stats(snap: Snapshot = Depends(current)) -> dict[str, Any]:
    return snap.meta["counts"]


@app.get("/api/graph/operations")
def graph_operations() -> dict[str, str]:
    return ops.OPERATIONS


@app.post("/api/graph/operations/{name}")
def run_named_operation(name: str, args: dict[str, str | None] | None = None, snap: Snapshot = Depends(current)) -> dict[str, Any]:
    if name not in ops.OPERATIONS:
        raise HTTPException(404, f"unknown operation {name}")
    try:
        return ops.run_operation(snap, name, args or {})
    except TypeError as error:
        raise HTTPException(422, f"bad arguments for {name}: {error}") from error


@app.get("/api/graph/gaps")
def graph_gaps(snap: Snapshot = Depends(current)) -> dict[str, Any]:
    return ops.gap_report(snap)


@app.get("/api/graph/viz/schema")
def viz_schema(snap: Snapshot = Depends(current)) -> dict[str, Any]:
    types = Counter(n["type"] for n in snap.nodes.values())
    issues = defaultdict(Counter)
    for issue in snap.issues:
        if issue["status"] == "open":
            issues[snap.nodes[issue["subject"]]["type"]][issue["severity"]] += 1
    links = Counter((snap.nodes[e["source"]]["type"], e["kind"], snap.nodes[e["target"]]["type"], e["method"]) for e in snap.edges)
    catalog_ = snap.catalog
    declared = {r.name for r in catalog_.relationship_types}
    gaps = [{"source": r.source, "kind": r.name, "target": r.target, "method": r.method, "count": 0} for r in catalog_.relationship_types if r.name not in {k for _, k, _, _ in links}]
    base_types = {t.name for t in room_catalog(None).entity_types}
    return {
        "types": [{"type": t, "count": types.get(t, 0), "issues": dict(issues.get(t, {})), "origin": "base" if t in base_types else "project"} for t in snap.type_order],
        "links": [{"source": s, "kind": k, "target": t, "method": m, "count": c} for (s, k, t, m), c in sorted(links.items())] + gaps,
        "declared": sorted(declared),
    }


@app.get("/api/graph/viz/search")
def viz_search(q: str = Query(min_length=1, max_length=100), limit: int = Query(30, le=100), snap: Snapshot = Depends(current)) -> list[dict[str, Any]]:
    text = q.strip().lower()
    digits = re.sub(r"\D", "", text)
    hits = []
    for node in snap.nodes.values():
        names = [node["label"].lower(), node["id"].lower(), *(a.lower() for a in snap.aliases.get(node["id"], []))]
        score = 100 if text in names else 70 if any(n.startswith(text) for n in names) else 40 if any(text in n for n in names) else 0
        if not score and len(digits) >= 6 and node["type"] == "Well" and digits in node["props"].get("api_14", ""):
            score = 80
        if score:
            hits.append((score + (1 if snap.issues_by_subject.get(node["id"]) else 0) - _type_index(snap, node["type"]) * 0.01, node))
    return [_node_card(snap, node) | {"score": round(score, 2)} for score, node in sorted(hits, key=lambda h: -h[0])[:limit]]


@app.get("/api/graph/viz/neighborhood")
def viz_neighborhood(id: str, depth: int = Query(2, ge=1, le=2), types: str | None = None, methods: str | None = None, issues_only: bool = False, snap: Snapshot = Depends(current)) -> dict[str, Any]:
    if id not in snap.nodes:
        raise HTTPException(404, f"unknown entity {id}")
    allowed_types = set(types.split(",")) if types else set(snap.type_order)
    allowed_methods = set(methods.split(",")) if methods else {"source", "key", "spatial", "alias", "document", "crosswalk", "rule", "human"}
    conflict = _conflict_edges(snap)
    visible = lambda e: e["method"] in allowed_methods or e["id"] in conflict  # noqa: E731
    level, parents = {id: 0}, {}
    frontier = [id]
    hub = len(snap.out.get(id, [])) + len(snap.inn.get(id, [])) > 60
    depth = 1 if hub else depth
    for d in range(1, depth + 1):
        nxt = []
        for current_ in frontier:
            for edge in [*snap.out.get(current_, []), *snap.inn.get(current_, [])]:
                other = edge["target"] if edge["source"] == current_ else edge["source"]
                if not visible(edge) or snap.nodes[other]["type"] not in allowed_types:
                    continue
                if other not in level:
                    level[other], parents[other] = d, [current_]
                    nxt.append(other)
                elif level[other] == d:
                    parents[other].append(current_)
        frontier = nxt
    caps = {1: 28, 2: 10}
    keep, capped = {id}, 0
    for d in (1, 2):
        groups = defaultdict(list)
        for node_id, l in level.items():
            if l == d:
                groups[snap.nodes[node_id]["type"]].append(node_id)
        for members in groups.values():
            has_issue = lambda n: bool(_issue_counts(snap, n))  # noqa: E731
            members.sort(key=lambda n: (not has_issue(n), -len(parents.get(n, [])), n))
            pool = [m for m in members if has_issue(m)] if issues_only and d == 2 else members
            chosen = [m for m in pool[: caps[d]] if d == 1 or any(p in keep for p in parents[m])]
            keep.update(chosen)
            capped += len(members) - len(chosen)
    edges = [e for e in snap.edges if e["source"] in keep and e["target"] in keep and visible(e)]
    return {
        "focus": id,
        "depth": depth,
        "hub": hub,
        "capped": capped,
        "nodes": [_node_card(snap, snap.nodes[n], level[n]) | {"parents": parents.get(n, [])} for n in sorted(keep, key=lambda n: (level[n], _type_index(snap, snap.nodes[n]["type"]), snap.nodes[n]["label"]))],
        "edges": [{"id": e["id"], "source": e["source"], "target": e["target"], "kind": e["kind"], "method": e["method"], "conflict": e["id"] in conflict, "evidence": e["locator"], "file": e["file"], "props": e["props"]} for e in edges],
    }


@app.get("/api/graph/viz/interpret")
def viz_interpret(id: str, snap: Snapshot = Depends(current)) -> dict[str, Any]:
    node = snap.nodes.get(id)
    if node is None:
        raise HTTPException(404, f"unknown entity {id}")
    conflict = _conflict_edges(snap)
    links = []
    for edge, direction in [*((e, "out") for e in snap.out.get(id, [])), *((e, "in") for e in snap.inn.get(id, []))]:
        other = snap.nodes[edge["target"] if direction == "out" else edge["source"]]
        links.append({"direction": direction, "kind": edge["kind"], "method": edge["method"], "file": edge["file"], "locator": edge["locator"], "props": edge["props"], "conflict": edge["id"] in conflict, "other": _node_card(snap, other)})
    facts = snap.facts(id)
    geometry = snap.sql("select kind, wkb from geometries where entity = ?", [id])
    from shapely import from_wkb
    from shapely.geometry import mapping as geo_mapping

    catalog_ = snap.catalog
    entity_type = next((t for t in catalog_.entity_types if t.name == node["type"]), None)
    return {
        "node": _node_card(snap, node) | {"file": node["file"], "locator": node["locator"], "props": node["props"]},
        "aliases": snap.aliases.get(id, []),
        "links": links[:400],
        "link_count": len(links),
        "facts": facts[:600],
        "fact_count": len(facts),
        "issues": snap.issues_by_subject.get(id, []),
        "geometry": [geo_mapping(from_wkb(row[1])) for row in geometry],
        "schema": entity_type.model_dump(mode="json") if entity_type else {"name": node["type"], "id_prefix": node["id"].split(":")[0], "key": "key", "description": "", "properties": []},
        "measures": [m.model_dump(mode="json") for m in catalog_.measures if m.subject == node["type"]],
    }


@app.post("/api/graph/ask")
async def graph_ask(request: AskRequest, snap: Snapshot = Depends(current)) -> dict[str, Any]:
    return await chat.ask(snap, request.question, request.scope)


# ---------------------------------------------------------------------------
# Files
# ---------------------------------------------------------------------------


@app.get("/api/files")
def files(snap: Snapshot = Depends(current)) -> list[dict[str, Any]]:
    by_file = defaultdict(Counter)
    for issue in snap.issues:
        for assertion in issue["assertions"]:
            if assertion.get("file"):
                by_file[assertion["file"]][issue["rule"]] += 1
    sheets = defaultdict(list)
    for sheet in snap.sheets:
        sheets[sheet["file_id"]].append(sheet)
    pending = Counter(m["file"] for m in snap.mappings if m["status"] == "proposed")
    proposals = {p["file"]: p for p in list_proposals(snap.room_id) if p["status"] == "proposed"}
    items = []
    for entry in snap.manifest["files"]:
        if entry["kind"] == "gis_sidecar":
            continue
        document = snap.documents.get(entry["file_id"])
        items.append({
            **{k: entry[k] for k in ("file_id", "path", "kind", "role", "size_bytes", "sha256", "as_of", "protection", "bundle")},
            "sidecars": sum(1 for f in snap.manifest["files"] if f["bundle"] == entry["bundle"] and f["kind"] == "gis_sidecar") if entry["bundle"] else 0,
            "sheets": [{k: s[k] for k in ("sheet", "rows", "columns", "filled_cells", "formula_cells", "uncached_formulas")} for s in sheets.get(entry["file_id"], [])],
            "pages": len(document["pages"]) if document else None,
            "figures": len(document["figures"]) if document else None,
            "claims": len(document["claims"]) if document else None,
            "issues": dict(by_file.get(entry["path"], {})),
            "incoming": entry["path"].startswith("incoming/"),
            "pending_mappings": pending.get(entry["path"], 0),
            "model_proposal": proposals[entry["path"]]["proposal_id"] if entry["path"] in proposals else None,
        })
    return items


@app.get("/api/files/{file_id}")
def file_detail(file_id: str, snap: Snapshot = Depends(current)) -> dict[str, Any]:
    entry = _entry(snap, file_id)
    datasets = [s for s in snap.schemas if s["file"] == entry["path"]]
    tables = {d: t for d, t in snap.tables.items() if t["file_id"] == file_id}
    mappings = [m for m in snap.mappings if m["file"] == entry["path"]]
    claims = {n["props"]["claim_id"]: {k: n["props"].get(k) for k in ("check", "computed", "computed_from")} | {"entity": n["id"]} for n in snap.nodes.values() if n["type"] == "DocumentClaim" and n["file"] == entry["path"]}
    proposals = [p for p in list_proposals(snap.room_id) if p["file"] == entry["path"]]
    return {"entry": entry, "datasets": datasets, "tables": tables, "mappings": mappings, "document": snap.documents.get(file_id), "claim_checks": claims, "proposals": proposals}


@app.get("/api/files/{file_id}/grid")
def file_grid(file_id: str, sheet: str, r0: int = Query(1, ge=1), r1: int = Query(60, ge=1), c0: int = Query(1, ge=1), c1: int = Query(30, ge=1), snap: Snapshot = Depends(current)) -> dict[str, Any]:
    grids = _grids(snap.room_id, snap.snapshot_id, file_id)
    grid = grids.get(sheet)
    if grid is None:
        raise HTTPException(404, f"unknown sheet {sheet}")
    r1, c1 = min(r1, r0 + 400, grid.max_row), min(c1, c0 + 80, grid.max_col)

    def cell(r: int, c: int) -> Any:
        value = grid.get(r, c)
        return value.isoformat()[:10] if hasattr(value, "isoformat") else value

    return {
        "sheet": sheet, "max_row": grid.max_row, "max_col": grid.max_col, "r0": r0, "r1": r1, "c0": c0, "c1": c1,
        "rows": [[cell(r, c) for c in range(c0, c1 + 1)] for r in range(r0, r1 + 1)],
        "formulas": [[c for c in range(c0, c1 + 1) if grid.formula(r, c)] for r in range(r0, r1 + 1)],
    }


@app.get("/api/files/{file_id}/lineage")
def cell_lineage(file_id: str, sheet: str, cell: str = Query(pattern=r"^[A-Z]{1,3}\d{1,6}$"), snap: Snapshot = Depends(current)) -> dict[str, Any]:
    path = _entry(snap, file_id)["path"]
    locator = f"{sheet}!{cell}"
    facts = snap.sql("select subject, measure, value, unit, zone, period, item, quality from facts where file = ? and locator = ?", [path, locator])
    edges = [e for e in snap.edges if e["file"] == path and e["locator"] == locator]
    nodes = [n for n in snap.nodes.values() if n["file"] == path and n["locator"] == locator]
    row = int(re.sub(r"\D", "", cell))
    entities = {f[0] for f in facts} | {e["source"] for e in edges} | {n["id"] for n in nodes}
    if not entities:
        row_pattern = re.compile(rf"^{re.escape(sheet)}![A-Z]+{row}$")
        entities = {n["id"] for n in snap.nodes.values() if n["file"] == path and n["locator"] and row_pattern.match(n["locator"])}
        if not entities:
            entities = {s for s, loc in snap.sql("select distinct subject, locator from facts where file = ? and locator like ?", [path, f"{sheet}!%{row}"]) if row_pattern.match(loc)}
    mapping = next((m for m in snap.mappings if m["file"] == path and m["column"] == re.sub(r"\d", "", cell) and snap.tables.get(m["dataset_id"], {}).get("sheet") == sheet), None)
    grid = _grids(snap.room_id, snap.snapshot_id, file_id).get(sheet) if _entry(snap, file_id)["kind"] in {"excel", "csv"} else None
    value = grid.get(row, column_index_from_string(re.sub(r"\d", "", cell))) if grid else None
    return {
        "locator": locator,
        "value": value.isoformat()[:10] if hasattr(value, "isoformat") else value,
        "facts": [dict(zip(["subject", "measure", "value", "unit", "zone", "period", "item", "quality"], f)) for f in facts],
        "links": [{"kind": e["kind"], "source": e["source"], "target": e["target"], "method": e["method"]} for e in edges],
        "entities": [_node_card(snap, snap.nodes[e]) for e in sorted(entities) if e in snap.nodes],
        "mapping": mapping,
        "issues": [i for e in entities for i in snap.issues_by_subject.get(e, []) if any(a.get("locator") == locator for a in i["assertions"])][:10],
    }


@app.get("/api/files/{file_id}/pages/{page}.png")
def page_image(file_id: str, page: int, snap: Snapshot = Depends(current)) -> FileResponse:
    path = snap.silver_dir / "documents" / file_id / "pages" / f"page-{page}.png"
    if file_id not in snap.documents or not path.exists():
        raise HTTPException(404, "page not found")
    return FileResponse(path, media_type="image/png")


@app.get("/api/files/{file_id}/figures/{figure_id}.png")
def figure_image(file_id: str, figure_id: str, snap: Snapshot = Depends(current)) -> FileResponse:
    if file_id not in snap.documents or not re.fullmatch(r"[\d.]+", figure_id):
        raise HTTPException(404, "figure not found")
    path = snap.silver_dir / "documents" / file_id / "figures" / f"figure-{figure_id}.png"
    if not path.exists():
        raise HTTPException(404, "figure not found")
    return FileResponse(path, media_type="image/png")


@app.get("/api/layers/{file_id}")
def layer(file_id: str, snap: Snapshot = Depends(current)) -> JSONResponse:
    _entry(snap, file_id)
    path = snap.silver_dir / "gis" / f"{file_id}.geojson"
    if not path.exists():
        raise HTTPException(404, "layer not found")
    data = json.loads(path.read_text("utf-8"))
    prefix = {"gis-pdp": "well", "gis-leasehold": "tract", "gis-units-pf": "unit", "gis-units": "unit"}.get(file_id)
    for feature in data["features"]:
        fid = str(feature["properties"]["feature_id"])
        entity = f"{prefix}:{fid[2:]}" if fid.startswith("g-") else f"{prefix}:{fid}" if prefix == "well" else f"{prefix}:{slug(fid)}"
        feature["properties"] = {"entity": entity if entity in snap.nodes else None, "issues": _issue_counts(snap, entity), **feature["properties"]}
    return JSONResponse(data)


# ---------------------------------------------------------------------------
# Review, mappings, rebuild, uploads, publish
# ---------------------------------------------------------------------------


@app.get("/api/issues")
def issues(rule: str | None = None, subject: str | None = None, status: str | None = None, issue_id: str | None = None, limit: int = Query(200, le=2000), offset: int = Query(0, ge=0), snap: Snapshot = Depends(current)) -> dict[str, Any]:
    selected = [i for i in snap.issues if (not rule or i["rule"] == rule) and (not subject or i["subject"] == subject) and (not status or i["status"] == status) and (not issue_id or i["issue_id"] == issue_id)]
    page = selected[offset: offset + limit]
    return {"total": len(selected), "items": [i | {"subject_label": snap.nodes[i["subject"]]["label"], "subject_type": snap.nodes[i["subject"]]["type"]} for i in page]}


@app.get("/api/rules")
def rules(snap: Snapshot = Depends(current)) -> list[dict[str, Any]]:
    counts = defaultdict(Counter)
    for issue in snap.issues:
        counts[issue["rule"]][issue["status"]] += 1
    return [{**r.model_dump(mode="json"), "open": counts[r.code]["open"], "decided": counts[r.code]["decided"], "policy": snap.policy(r.code)} for r in load_rules().rules if counts[r.code]]


@app.post("/api/issues/decision")
def issue_decision(decision: IssueDecision, snap: Snapshot = Depends(current)) -> dict[str, Any]:
    if decision.scope == "issue" and decision.issue_id not in snap.issue_index:
        raise HTTPException(404, "unknown issue")
    record = record_issue_decision(snap.room_id, decision.issue_id or f"rule:{decision.rule}", decision.rule, decision.choice, decision.reason, decision.reviewer, snap.snapshot_id, decision.scope)
    snap.refresh_decisions()
    return record


@app.post("/api/publish")
def publish_snapshot(request: PublishRequest, room: str = Depends(room_param)) -> dict[str, Any]:
    state.snapshot(room)
    outcome = publish(room, reviewer=request.reviewer, reason=request.reason, force=request.force)
    state.reload(room)
    return outcome


@app.get("/api/mappings")
def mappings(status: str | None = None, file_id: str | None = None, snap: Snapshot = Depends(current)) -> list[dict[str, Any]]:
    path = _entry(snap, file_id)["path"] if file_id else None
    return [m for m in snap.mappings if (not status or m["status"] == status) and (not path or m["file"] == path)]


@app.post("/api/mappings/decision")
def mapping_decision(decision: MappingDecision, snap: Snapshot = Depends(current)) -> dict[str, Any]:
    proposal = next((m for m in snap.mappings if m["field_id"] == decision.field_id and m["target"] == decision.target), None)
    if proposal is None:
        raise HTTPException(404, "unknown mapping proposal")
    record = record_mapping_decision(snap.room_id, {k: proposal[k] for k in ("field_id", "header", "target", "kind", "method")} | {"header_key": header_key(proposal["header"])}, decision.status, decision.reviewer, decision.reason)
    proposal.update({"status": decision.status, "decided_by": decision.reviewer})
    return record


@app.post("/api/rebuild")
async def rebuild(room: str = Depends(room_param)) -> dict[str, Any]:
    from ..azure_services import azure_options, mapping_model
    from ..ontology.runner import BuildRequest, ReviewResponse, run_ontology_workflow

    record = get_project(room)
    if record is None:
        raise HTTPException(404, f"unknown project {room}")
    result = await run_ontology_workflow(BuildRequest(room, str(room_dir(record))), {**azure_options(), "mapping_model": mapping_model()}, ReviewResponse("hold"))
    snap = state.reload(room)
    return {
        "snapshot_id": result.snapshot_id, "status": snap.meta["status"], "counts": snap.meta["counts"], "mappings": snap.meta.get("mappings"),
        "catalog_version": snap.meta.get("catalog_version"), "proposals": [p for p in list_proposals(room) if p["status"] == "proposed"],
    }


def _safe_name(name: str) -> str:
    return re.sub(r"[^A-Za-z0-9 ._()-]", "_", name)


def _unpack(content: bytes, incoming: Path) -> list[str]:
    """Extract a zipped data room (or shapefile bundle) into incoming/, keeping folders, refusing unsafe paths."""
    written = []
    with zipfile.ZipFile(io.BytesIO(content)) as archive:
        for member in archive.infolist():
            name = Path(member.filename)
            if member.is_dir() or any(part.startswith((".", "~$", "__MACOSX")) for part in name.parts) or name.suffix.lower() not in ZIP_MEMBER_TYPES:
                continue
            if name.is_absolute() or ".." in name.parts:
                raise HTTPException(400, f"unsafe path in archive: {member.filename}")
            target = incoming.joinpath(*(_safe_name(part) for part in name.parts))
            target.parent.mkdir(parents=True, exist_ok=True)
            target.write_bytes(archive.read(member))
            written.append(target.relative_to(incoming).as_posix())
    if not written:
        raise HTTPException(400, "the archive contains no workbook, CSV, PDF or shapefile")
    return written


@app.post("/api/uploads")
async def upload(file: UploadFile = File(...), room: str = Depends(room_param)) -> dict[str, Any]:
    record = get_project(room)
    if record is None:
        raise HTTPException(404, f"unknown project {room}")
    name = Path(file.filename or "").name
    if not name or Path(name).suffix.lower() not in UPLOAD_TYPES or name.startswith((".", "~$")):
        raise HTTPException(400, f"upload a {', '.join(sorted(UPLOAD_TYPES))} file")
    content = await file.read(MAX_UPLOAD_BYTES + 1)
    if len(content) > MAX_UPLOAD_BYTES:
        raise HTTPException(413, "file is larger than 200 MB")
    incoming = RoomPaths(room).incoming
    incoming.mkdir(parents=True, exist_ok=True)
    if name.lower().endswith(".zip"):
        written = _unpack(content, incoming)
    else:
        target = incoming / _safe_name(name)
        target.write_bytes(content)
        written = [target.name]
    rebuilt = await rebuild(room)
    snap = state.snapshot(room)
    entries = [f for f in snap.manifest["files"] if f["path"] in {f"incoming/{w}" for w in written} and f["kind"] != "gis_sidecar"]
    entry = entries[0] if entries else None
    return {
        **rebuilt,
        "file_id": entry["file_id"] if entry else None,
        "files": [e["file_id"] for e in entries],
        "proposals": [m for m in snap.mappings if entry and m["file"] == entry["path"]],
        "model_proposals": [p for p in list_proposals(room) if p["status"] == "proposed" and any(p["file"] == e["path"] for e in entries)],
    }


# ---------------------------------------------------------------------------
# Snapshots, diffs and exports
# ---------------------------------------------------------------------------


@app.get("/api/snapshots")
def snapshots(room: str = Depends(room_param)) -> list[dict[str, Any]]:
    if get_project(room) is None:
        raise HTTPException(404, f"unknown project {room}")
    return list_snapshots(room)


@app.get("/api/snapshots/diff")
def snapshots_diff(to: str | None = Query(None, alias="to"), from_: str | None = Query(None, alias="from"), room: str = Depends(room_param)) -> dict[str, Any]:
    """Changes between two snapshots; by default the latest against the one built before it."""
    history = list_snapshots(room)
    if not history:
        raise HTTPException(404, "no snapshots")
    ids = [s["snapshot_id"] for s in history]
    after = to or ids[-1]
    before = from_ or (ids[ids.index(after) - 1] if after in ids and ids.index(after) > 0 else None)
    if before is None:
        return {"before": None, "after": {"snapshot_id": after}, "message": "This is the first snapshot of the project."}
    if before not in ids or after not in ids:
        raise HTTPException(404, "unknown snapshot")
    return snapshot_diff(room, before, after)


def _csv(rows: list[dict[str, Any]], filename: str) -> StreamingResponse:
    buffer = io.StringIO()
    if rows:
        writer = csv.DictWriter(buffer, fieldnames=list(rows[0].keys()))
        writer.writeheader()
        writer.writerows(rows)
    buffer.seek(0)
    return StreamingResponse(iter([buffer.getvalue()]), media_type="text/csv", headers={"Content-Disposition": f'attachment; filename="{filename}"'})


@app.get("/api/export/decisions.csv")
def export_decisions(snap: Snapshot = Depends(current)) -> StreamingResponse:
    rows = []
    for record in issue_decisions(snap.room_id):
        issue = snap.issue_index.get(record["issue_id"])
        rows.append({**record, "subject": issue["subject"] if issue else "", "detail": issue["detail"] if issue else ""})
    return _csv(rows, f"{snap.room_id}-decisions.csv")


@app.get("/api/export/facts.csv")
def export_facts(snap: Snapshot = Depends(current)) -> StreamingResponse:
    rows = snap.facts()
    labels = {n["id"]: n["label"] for n in snap.nodes.values()}
    return _csv([{"subject": f["subject"], "subject_label": labels.get(f["subject"], ""), **{k: f[k] for k in ("measure", "value", "unit", "zone", "period", "item", "section", "file", "locator", "quality")}} for f in rows], f"{snap.room_id}-facts.csv")


@app.get("/api/export/entities.csv")
def export_entities(snap: Snapshot = Depends(current)) -> StreamingResponse:
    rows = [{"id": n["id"], "type": n["type"], "label": n["label"], "file": n["file"] or "", "locator": n["locator"] or "", "properties": json.dumps(n["props"], default=str)} for n in snap.nodes.values()]
    return _csv(rows, f"{snap.room_id}-entities.csv")


FRONTEND = REPO_ROOT / "frontend" / "dist"
if FRONTEND.exists():
    app.mount("/assets", StaticFiles(directory=FRONTEND / "assets"), name="assets")

    @app.get("/{path:path}", include_in_schema=False)
    def spa(path: str) -> FileResponse:
        if path.startswith("api/"):
            raise HTTPException(404, "unknown endpoint")
        candidate = (FRONTEND / path).resolve()
        if path and candidate.is_file() and candidate.is_relative_to(FRONTEND.resolve()):
            return FileResponse(candidate)
        return FileResponse(FRONTEND / "index.html")
